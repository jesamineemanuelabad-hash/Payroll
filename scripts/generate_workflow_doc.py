from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "workflow-diagrams"
OUTPUT = DOCS / "Payroll_Benefits_System_Workflow.docx"

NAVY = "0F172A"
SLATE = "475569"
MUTED = "64748B"
LIGHT = "F8FAFC"
BORDER = "CBD5E1"
INDIGO = "4F46E5"
INDIGO_LIGHT = "EEF2FF"
EMERALD = "059669"
EMERALD_LIGHT = "ECFDF5"
AMBER = "D97706"
AMBER_LIGHT = "FFFBEB"
RED = "DC2626"
RED_LIGHT = "FEF2F2"
BLUE = "0284C7"
BLUE_LIGHT = "F0F9FF"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


F_TITLE = font(50, True)
F_LANE = font(24, True)
F_BOX = font(27, True)
F_BODY = font(22)
F_SMALL = font(18)


def hex_color(value: str):
    value = value.lstrip("#")
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def rounded_box(draw, rect, title, body="", fill=INDIGO_LIGHT, outline=INDIGO, title_color=NAVY, body_color=SLATE):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=22, fill=hex_color(fill), outline=hex_color(outline), width=3)
    max_chars = max(16, int((x2 - x1) / 17))
    title_lines = wrap(title, max_chars)
    body_lines = wrap(body, max_chars + 8) if body else []
    total = len(title_lines) * 34 + (12 if body_lines else 0) + len(body_lines) * 28
    y = y1 + max(16, ((y2 - y1) - total) / 2)
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=F_BOX)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=F_BOX, fill=hex_color(title_color))
        y += 34
    if body_lines:
        y += 8
        for line in body_lines:
            bbox = draw.textbbox((0, 0), line, font=F_BODY)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=F_BODY, fill=hex_color(body_color))
            y += 28


def arrow(draw, start, end, color=SLATE, label=None):
    draw.line([start, end], fill=hex_color(color), width=5)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
    draw.polygon(points, fill=hex_color(color))
    if label:
        bbox = draw.textbbox((0, 0), label, font=F_SMALL)
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        draw.rounded_rectangle((mx - (bbox[2] - bbox[0]) / 2 - 8, my - 14, mx + (bbox[2] - bbox[0]) / 2 + 8, my + 14), 7, fill="white")
        draw.text((mx - (bbox[2] - bbox[0]) / 2, my - 11), label, font=F_SMALL, fill=hex_color(MUTED))


def canvas(title, width=2200, height=1250):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((20, 20, width - 20, height - 20), 28, fill="white", outline=hex_color(BORDER), width=3)
    draw.text((70, 52), title, font=F_TITLE, fill=hex_color(NAVY))
    draw.line((70, 125, width - 70, 125), fill=hex_color(BORDER), width=3)
    return image, draw


def save_end_to_end():
    image, draw = canvas("End-to-end operational workflow", height=1340)
    steps = [
        ("1. Access", "Password + TOTP MFA\nRBAC and active-account check", BLUE_LIGHT, BLUE),
        ("2. Prepare people data", "Create/import employees\nAssign department and compensation", INDIGO_LIGHT, INDIGO),
        ("3. Capture inputs", "Attendance, leave, benefits, claims,\ncompensation changes and bonuses", EMERALD_LIGHT, EMERALD),
        ("4. Calculate payroll", "Apply effective policy, cutoff rules,\nstatutory deductions and premiums", AMBER_LIGHT, AMBER),
        ("5. Review and approve", "Validate exceptions and comparisons\nSubmit -> approve -> mark paid", INDIGO_LIGHT, INDIGO),
        ("6. Deliver outputs", "Excel register, print/PDF payslips,\nlive dashboard and audit history", EMERALD_LIGHT, EMERALD),
    ]
    x_positions = [110, 760, 1410]
    y_positions = [190, 690]
    boxes = []
    for index, (title, body, fill, outline) in enumerate(steps):
        row, col = divmod(index, 3)
        x = x_positions[col] if row == 0 else x_positions[2 - col]
        y = y_positions[row]
        rect = (x, y, x + 560, y + 290)
        rounded_box(draw, rect, title, body, fill, outline)
        boxes.append(rect)
    arrow(draw, (670, 335), (760, 335))
    arrow(draw, (1320, 335), (1410, 335))
    arrow(draw, (1690, 480), (1690, 690))
    arrow(draw, (1410, 835), (1320, 835))
    arrow(draw, (760, 835), (670, 835))
    draw.rounded_rectangle((110, 1080, 1970, 1240), 20, fill=hex_color(LIGHT), outline=hex_color(BORDER), width=2)
    draw.text((155, 1110), "CONTROL PRINCIPLE", font=F_LANE, fill=hex_color(INDIGO))
    control = "Every protected operation is checked by Supabase authentication, database RBAC, row-level security, validation rules and an audit trail."
    for line_no, line in enumerate(wrap(control, 125)):
        draw.text((155, 1150 + line_no * 30), line, font=F_BODY, fill=hex_color(SLATE))
    path = ASSETS / "01-end-to-end-workflow.png"
    image.save(path, quality=95)
    return path


def save_access():
    image, draw = canvas("Authentication, MFA and authorization flow", height=1400)
    xs = [120, 750, 1380]
    rounded_box(draw, (xs[0], 180, xs[0] + 500, 390), "User signs in", "Supabase email and password", BLUE_LIGHT, BLUE)
    rounded_box(draw, (xs[1], 180, xs[1] + 500, 390), "Workspace session", "Active profile + assigned role required", INDIGO_LIGHT, INDIGO)
    rounded_box(draw, (xs[2], 180, xs[2] + 500, 390), "Privileged role?", "Super admin, HR, payroll or HR manager", AMBER_LIGHT, AMBER)
    arrow(draw, (620, 285), (750, 285))
    arrow(draw, (1250, 285), (1380, 285))
    rounded_box(draw, (1380, 510, 1880, 720), "TOTP verification", "AAL2 session required by UI and database", RED_LIGHT, RED)
    arrow(draw, (1630, 390), (1630, 510), label="Yes")
    rounded_box(draw, (750, 820, 1250, 1030), "Database authorization", "RBAC helper + RLS + employee scope", INDIGO_LIGHT, INDIGO)
    arrow(draw, (1380, 615), (1250, 925))
    arrow(draw, (1630, 390), (1120, 820), label="No")
    rounded_box(draw, (120, 820, 620, 1030), "Requested operation", "Read, create, edit, approve, export or delete", BLUE_LIGHT, BLUE)
    arrow(draw, (620, 925), (750, 925))
    rounded_box(draw, (1380, 820, 1880, 1030), "Decision", "Permit scoped operation or deny safely", EMERALD_LIGHT, EMERALD)
    arrow(draw, (1250, 925), (1380, 925))
    rounded_box(draw, (750, 1120, 1250, 1300), "Audit log", "Authorized record changes preserve actor, time and action", LIGHT, SLATE)
    arrow(draw, (1630, 1030), (1250, 1210))
    path = ASSETS / "02-security-access-flow.png"
    image.save(path, quality=95)
    return path


def save_attendance():
    image, draw = canvas("Employee, attendance and paid-leave workflow", height=1420)
    lanes = [(150, "ESS / HR input", BLUE_LIGHT), (510, "System validation", INDIGO_LIGHT), (870, "Payroll-ready output", EMERALD_LIGHT)]
    for y, label, fill in lanes:
        draw.rounded_rectangle((70, y, 2130, y + 285), 18, fill=hex_color(fill), outline=hex_color(BORDER), width=2)
        draw.text((95, y + 18), label, font=F_LANE, fill=hex_color(NAVY))
    top = [
        (150, "Employee/profile", "Identity, department, status"),
        (790, "Time records", "Time-in, time-out and work-day type"),
        (1430, "Leave request", "Draft -> submit -> HR decision"),
    ]
    for x, title, body in top:
        rounded_box(draw, (x, 225, x + 480, 395), title, body, "FFFFFF", BLUE)
    mid = [
        (150, "Normalize", "Idempotent sync and effective dates"),
        (790, "Compute minutes", "Worked, late, undertime, absence and overtime"),
        (1430, "Approve paid leave", "Self-approval blocked; weekdays counted"),
    ]
    for x, title, body in mid:
        rounded_box(draw, (x, 585, x + 480, 755), title, body, "FFFFFF", INDIGO)
    bottom = [
        (150, "Eligible worker", "Active/on-leave and payroll eligible"),
        (790, "Attendance totals", "Period minutes and work-day classifications"),
        (1430, "Paid working day", "Prevents absence deduction in payroll"),
    ]
    for x, title, body in bottom:
        rounded_box(draw, (x, 945, x + 480, 1115), title, body, "FFFFFF", EMERALD)
    for x in (390, 1030, 1670):
        arrow(draw, (x, 395), (x, 585))
        arrow(draw, (x, 755), (x, 945))
    draw.rounded_rectangle((390, 1220, 1810, 1340), 18, fill=hex_color(AMBER_LIGHT), outline=hex_color(AMBER), width=3)
    text = "Payroll calculation combines all three verified outputs for the selected cutoff period."
    bbox = draw.textbbox((0, 0), text, font=F_BOX)
    draw.text(((2200 - (bbox[2] - bbox[0])) / 2, 1261), text, font=F_BOX, fill=hex_color(NAVY))
    for x in (390, 1030, 1670):
        arrow(draw, (x, 1115), (1100, 1220))
    path = ASSETS / "03-attendance-leave-flow.png"
    image.save(path, quality=95)
    return path


def save_payroll():
    image, draw = canvas("Payroll calculation and approval workflow", height=1850)
    stages = [
        ("Create draft run", "1st-15th, 16th-month end, or monthly; preparation date is payday minus 2 days", BLUE_LIGHT, BLUE),
        ("Load eligible employees", "Active payroll employees with effective compensation; system owner excluded", INDIGO_LIGHT, INDIGO),
        ("Calculate period earnings", "Monthly/daily/hourly basic pay + allowances + approved bonus + reimbursements", EMERALD_LIGHT, EMERALD),
        ("Apply time adjustments", "Late + undertime + absence deductions; overtime and night differential premiums", AMBER_LIGHT, AMBER),
        ("Apply policy and deductions", "Effective SSS, PhilHealth, Pag-IBIG, BIR tax, benefits and authorized deductions", RED_LIGHT, RED),
        ("Recompute totals", "Gross pay - itemized deductions = net pay; employer contributions stored separately", INDIGO_LIGHT, INDIGO),
        ("Validate payroll", "Blocking exceptions, negative values, missing inputs and optional historical comparisons", AMBER_LIGHT, AMBER),
        ("Submit and approve", "Draft -> pending approval -> approved; self-approval and invalid transitions blocked", BLUE_LIGHT, BLUE),
        ("Pay and deliver", "Mark paid; export Excel register; print/PDF register and employee payslips", EMERALD_LIGHT, EMERALD),
    ]
    y = 165
    for idx, (title, body, fill, outline) in enumerate(stages):
        x = 180 if idx % 2 == 0 else 1120
        rect = (x, y, x + 820, y + 230)
        rounded_box(draw, rect, f"{idx + 1}. {title}", body, fill, outline)
        if idx < len(stages) - 1:
            if idx % 2 == 0:
                arrow(draw, (x + 820, y + 115), (1120, y + 115))
            else:
                arrow(draw, (1530, y + 230), (590, y + 335))
                y += 335
    draw.rounded_rectangle((1120, 1505, 1940, 1625), 18, fill=hex_color(LIGHT), outline=hex_color(BORDER), width=2)
    draw.text((1170, 1530), "IMMUTABILITY CONTROL", font=F_LANE, fill=hex_color(INDIGO))
    draw.text((1170, 1572), "Approved/non-draft payroll entries cannot be edited.", font=F_BODY, fill=hex_color(SLATE))
    path = ASSETS / "04-payroll-workflow.png"
    image.save(path, quality=95)
    return path


def save_analytics():
    image, draw = canvas("Automatic HR Analytics and XGBoost scoring flow", height=1450)
    nodes = [
        ((120, 180, 620, 390), "User opens HR Analytics", "Authorized super admin or HR admin with MFA", BLUE_LIGHT, BLUE),
        ((850, 180, 1350, 390), "Check scoring state", "Last 30 days changed or model older than one day?", INDIGO_LIGHT, INDIGO),
        ((1580, 180, 2080, 390), "Claim singleton lease", "Stops concurrent visits from starting duplicate jobs", AMBER_LIGHT, AMBER),
        ((1580, 560, 2080, 770), "Build feature payload", "Minutes, rolling late/absence rates, day of week and history", INDIGO_LIGHT, INDIGO),
        ((850, 560, 1350, 770), "XGBoost service", "Classifications, probabilities, anomaly score and reasons", RED_LIGHT, RED),
        ((120, 560, 620, 770), "Validate and persist", "Versioned model run and predictions saved in PostgreSQL", EMERALD_LIGHT, EMERALD),
        ((120, 960, 620, 1170), "Dashboard insights", "Attendance patterns, anomalies and model accuracy", BLUE_LIGHT, BLUE),
        ((850, 960, 1350, 1170), "Human review", "Predictions are review signals; they never change payroll directly", AMBER_LIGHT, AMBER),
        ((1580, 960, 2080, 1170), "Export reporting", "Filtered live analytics to CSV or Excel", EMERALD_LIGHT, EMERALD),
    ]
    for rect, title, body, fill, outline in nodes:
        rounded_box(draw, rect, title, body, fill, outline)
    arrow(draw, (620, 285), (850, 285))
    arrow(draw, (1350, 285), (1580, 285), label="Stale")
    arrow(draw, (1830, 390), (1830, 560))
    arrow(draw, (1580, 665), (1350, 665))
    arrow(draw, (850, 665), (620, 665))
    arrow(draw, (370, 770), (370, 960))
    arrow(draw, (620, 1065), (850, 1065))
    arrow(draw, (1350, 1065), (1580, 1065))
    draw.rounded_rectangle((850, 1240, 2080, 1350), 18, fill=hex_color(LIGHT), outline=hex_color(BORDER), width=2)
    draw.text((900, 1267), "If current: load existing live snapshot. If service fails: release lease and retry after 15 minutes.", font=F_BODY, fill=hex_color(SLATE))
    path = ASSETS / "05-analytics-flow.png"
    image.save(path, quality=95)
    return path


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, True, "FFFFFF", 9)
        set_cell_shading(table.rows[0].cells[idx], INDIGO)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, value in enumerate(values):
            set_cell_text(cells[col_idx], value, False, NAVY, 8.5)
            if row_idx % 2:
                set_cell_shading(cells[col_idx], LIGHT)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    document.add_paragraph()
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_callout(document, title, body, color=INDIGO_LIGHT):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r.font.size = Pt(10)
    r2 = p.add_run(body)
    r2.font.color.rgb = RGBColor.from_string(SLATE)
    r2.font.size = Pt(9)
    document.add_paragraph()


def add_diagram(document, path, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.85))
    caption_p = document.add_paragraph(caption)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.style = document.styles["Caption"]


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(SLATE)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for level, size, color in [(1, 20, NAVY), (2, 14, INDIGO), (3, 11, NAVY)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    header = section.header.paragraphs[0]
    header.text = "PRIORITY HANDLING LOGISTICS, INC.  |  PAYROLL & BENEFITS"
    header.style = styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Internal system workflow  |  Generated 19 September 2026  |  Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def page_break(document):
    document.add_page_break()


def build_document(diagrams):
    document = Document()
    configure_document(document)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(80)
    mark = cover.add_run("PAYROLL & BENEFITS")
    mark.bold = True
    mark.font.name = "Aptos"
    mark.font.size = Pt(12)
    mark.font.color.rgb = RGBColor.from_string(INDIGO)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("System Workflow &\nOperational Flowcharts")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Priority Handling Logistics, Inc.\nImplementation-aligned process guide").font.size = Pt(14)
    document.add_paragraph("\n")
    add_callout(document, "Document purpose", "A client-ready explanation of how users, attendance, leave, compensation, benefits, claims, payroll, reports and AI-assisted attendance analytics move through the implemented system.")
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("Version 1.0  |  19 September 2026\nPrepared from the current Next.js, Supabase and PostgreSQL implementation").font.size = Pt(9)

    page_break(document)
    add_heading(document, "Document control", 1)
    add_table(document, ["Item", "Value"], [
        ("Document", "Payroll & Benefits System Workflow & Operational Flowcharts"),
        ("Organization", "Priority Handling Logistics, Inc."),
        ("Audience", "Super administrators, HR administrators, payroll managers, HR managers, reviewers and implementation teams"),
        ("Source of truth", "Current application routes, server actions and ordered Supabase migrations through 202609190001"),
        ("Status", "Implementation-aligned working document; statutory and company policy values require owner approval before production payroll"),
    ], [1.5, 5.1])
    add_heading(document, "Contents", 1)
    add_numbered(document, [
        "Executive overview and system boundaries",
        "Roles, access and security workflow",
        "Employee, attendance and leave workflow",
        "Payroll calculation and approval workflow",
        "Compensation, benefits and claims workflow",
        "Automatic HR Analytics workflow",
        "Reports, exports, payslips and auditability",
        "Recommended operating procedure and exception handling",
        "Deployment and acceptance checklist",
    ])
    add_callout(document, "Reading guide", "Green paths represent completed, payroll-ready outputs. Amber paths require review. Red controls represent security, validation or exception gates. Dashed external boundaries are described in text because they require a configured third-party service.", AMBER_LIGHT)

    page_break(document)
    add_heading(document, "1. Executive overview", 1)
    document.add_paragraph("The Payroll & Benefits system centralizes employee records, attendance, paid leave, compensation, benefits, claims, payroll calculation, payslips and analytics. PostgreSQL is the authoritative operational store. Next.js supplies the responsive user interface and server actions, while Supabase provides authentication, MFA, database access and row-level security.")
    add_diagram(document, diagrams[0], "Figure 1. End-to-end operational workflow from secure access to auditable outputs.")
    add_heading(document, "1.1 Implemented automation", 2)
    add_bullets(document, [
        "Validated CRUD, search, pagination, confirmed deletion and audit history for operational records.",
        "Effective-dated employee compensation and effective-dated payroll policy selection.",
        "1st cutoff (1-15), 2nd cutoff (16-month end) and full-month payroll schedules.",
        "Late, undertime, absence, overtime, night differential, approved paid leave, bonuses and reimbursements.",
        "Itemized SSS, PhilHealth, Pag-IBIG and BIR withholding deductions through a versioned policy.",
        "Payroll validation, controlled state transitions, immutable approved payroll, Excel registers and print/PDF payslips.",
        "Live Overview and HR Analytics snapshots; automatic on-visit XGBoost scoring when configured.",
    ])
    add_heading(document, "1.2 External or policy-dependent boundaries", 2)
    add_bullets(document, [
        "ESS synchronization needs a configured server-side ESS integration endpoint.",
        "XGBoost scoring needs XGBOOST_SERVICE_URL and XGBOOST_SERVICE_TOKEN plus the documented prediction contract.",
        "Bank disbursement, file uploads, loan deductions, 13th-month pay and automatic government holiday synchronization are separate integrations.",
        "Government rates, taxable benefits, cutoff allocation, schedules and rounding must be approved and updated through reviewed policy versions.",
    ])

    page_break(document)
    add_heading(document, "2. Roles, access and security workflow", 1)
    add_diagram(document, diagrams[1], "Figure 2. Authentication, MFA, RBAC, row-level security and audit controls.")
    add_table(document, ["Role", "Primary responsibility", "Typical access"], [
        ("super_admin", "System ownership and governance", "All modules; invitations; RBAC; payroll policy; protected owner safeguards"),
        ("hr_admin", "People and HR operations", "Employees, attendance, leave, compensation, benefits, claims and analytics"),
        ("payroll_manager", "Payroll preparation and processing", "Attendance inputs, payroll calculation, validation, approval workflow and outputs"),
        ("hr_manager", "HR review", "Compensation and claims review; permitted analytics"),
        ("manager", "Department oversight", "Read only within assigned department scope"),
        ("employee", "Self-service visibility", "Own permitted records; no administrative CRUD"),
    ], [1.1, 2.2, 3.3])
    add_heading(document, "2.1 Security controls", 2)
    add_bullets(document, [
        "Privileged roles require TOTP MFA and an AAL2 token. Database role checks enforce MFA, not only the interface.",
        "The system owner remains active, cannot lose the super-admin role and is excluded from payroll employee records.",
        "Inactive or terminated profiles are denied workspace access immediately.",
        "Self-approval is blocked for governed leave, claims and compensation decisions.",
        "Direct authenticated writes to sensitive role and workflow tables are revoked; validated RPCs perform atomic changes.",
    ])

    page_break(document)
    add_heading(document, "3. Employee, attendance and leave workflow", 1)
    add_diagram(document, diagrams[2], "Figure 3. Employee identity, attendance minutes and approved paid leave converge into payroll-ready inputs.")
    add_heading(document, "3.1 Employee preparation", 2)
    add_numbered(document, [
        "Create or invite the Supabase Authentication account.",
        "Create the employee profile, employee number, department, job title, location and employment status.",
        "Mark the person as payroll eligible; system identities remain excluded.",
        "Create an effective compensation history record with salary amount and frequency.",
        "Assign the appropriate workspace role and optional manager department scope.",
    ])
    add_heading(document, "3.2 Attendance inputs", 2)
    add_table(document, ["Input", "System treatment", "Payroll effect"], [
        ("Time-in / time-out", "Stored as timestamps and validated for logical order", "Determines worked minutes"),
        ("Late minutes", "Summed within the payroll period", "Hourly rate x late minutes / 60"),
        ("Undertime minutes", "Summed independently; does not offset overtime", "Hourly rate x undertime minutes / 60"),
        ("Absence", "Recorded classification and absence minutes", "Fixed-salary absence deduction unless covered by paid leave"),
        ("Overtime minutes", "Uses reviewed work-day type", "Policy multiplier by ordinary/rest/special/holiday category"),
        ("Night minutes", "Stored for eligible hours", "Policy night differential"),
    ], [1.4, 2.7, 2.5])
    add_heading(document, "3.3 Paid-leave process", 2)
    add_numbered(document, [
        "Employee/HR creates a draft leave request and submits it.",
        "An authorized HR reviewer approves or rejects the submitted request; the requester cannot approve their own request.",
        "Approved paid weekdays overlapping the payroll period are counted as paid leave days.",
        "Paid leave is included as a paid working day and prevents the related absence deduction.",
        "Decided requests are protected from casual edits or deletion.",
    ])

    page_break(document)
    add_heading(document, "4. Payroll calculation and approval workflow", 1)
    add_diagram(document, diagrams[3], "Figure 4. Calculation, validation, approval, payment and output lifecycle.")
    add_heading(document, "4.1 Core computation", 2)
    add_table(document, ["Component", "Implemented calculation"], [
        ("Monthly / semi-monthly basic", "Applicable period share based on the selected cutoff schedule"),
        ("Daily-paid basic", "Daily rate x (worked days + approved paid-leave days)"),
        ("Hourly-paid basic", "Hourly rate x (worked hours + approved paid-leave hours)"),
        ("Late / undertime", "Hourly rate x minutes / 60"),
        ("Fixed-salary absence", "Daily rate x absence minutes / configured workday minutes"),
        ("Overtime", "Hourly rate x eligible hours x effective day-type multiplier"),
        ("Night differential", "Hourly rate x night hours x configured differential"),
        ("Gross pay", "Basic + allowances + overtime + night differential + bonus + benefits + reimbursements"),
        ("Net pay", "Gross pay - itemized employee deductions"),
    ], [1.8, 4.8])
    add_heading(document, "4.2 Statutory and other deductions", 2)
    document.add_paragraph("The effective payroll policy controls statutory rates, contribution caps, cutoff allocation, premium multipliers and rounding. The seeded policy reflects the configured 2025/2023 rule references but must be reviewed by the company before live payroll.")
    add_bullets(document, [
        "SSS employee and employer contributions, including employer EC treatment.",
        "PhilHealth employee and employer shares with configured salary floor and ceiling.",
        "Pag-IBIG employee and employer shares with configured compensation cap.",
        "BIR withholding tax for semi-monthly and monthly schedules.",
        "Benefit employee cost, authorized other deductions and non-taxable approved expense reimbursements.",
    ])
    add_heading(document, "4.3 Payroll state controls", 2)
    add_table(document, ["State", "Allowed activity", "Control"], [
        ("Draft", "Edit inputs, calculate, recalculate, validate", "Only draft entries may change"),
        ("Processing / needs review", "Resolve calculation exceptions", "Blocking validation issues must be cleared"),
        ("Pending approval", "Authorized approver reviews totals", "Invalid/self approval is blocked"),
        ("Approved", "Generate final outputs and prepare payment", "Lock timestamp; entries immutable"),
        ("Paid", "Historical reporting and employee payslip access", "Final state retained in audit history"),
    ], [1.2, 2.6, 2.8])

    page_break(document)
    add_heading(document, "5. Compensation, benefits and claims workflow", 1)
    add_heading(document, "5.1 Compensation changes", 2)
    add_numbered(document, [
        "Create a compensation cycle and a draft proposal containing current salary, proposed salary, bonus and justification.",
        "Submit the proposal for review. A different authorized reviewer approves or rejects it.",
        "Approval creates the effective salary-history change transactionally and preserves compensation history.",
        "Payroll loads the salary record effective for the payroll period; an approved bonus is applied idempotently.",
    ])
    add_heading(document, "5.2 Benefits", 2)
    add_numbered(document, [
        "Maintain providers and plans for HMO, allowances, insurance/group, leave and government categories.",
        "Enroll eligible employees with effective/expiration dates and membership details.",
        "During calculation, active eligible plans contribute the configured employee cost and employer cost.",
        "Benefit history, record exports and audit events remain available to permitted roles.",
    ])
    add_heading(document, "5.3 Claims and reimbursements", 2)
    add_numbered(document, [
        "Create a draft claim with category, amount, description and HTTPS supporting-document link.",
        "Submit the claim; verify document status; approve or reject with the required reason.",
        "Approved claims not already linked to a payroll are included as reimbursements during calculation.",
        "The claim is linked to the payroll run to prevent duplicate reimbursement.",
    ])
    add_callout(document, "Review boundary", "The system records a document URL and verification decision. It does not currently upload or virus-scan files; production document storage requires a separate secure integration.", AMBER_LIGHT)

    page_break(document)
    add_heading(document, "6. Automatic HR Analytics workflow", 1)
    add_diagram(document, diagrams[4], "Figure 5. Automatic, on-visit XGBoost attendance scoring with concurrency and retry controls.")
    add_heading(document, "6.1 Automatic behavior", 2)
    add_numbered(document, [
        "An authorized user opens HR Analytics; no separate Run AI button is required.",
        "The database checks whether the last 30 days of attendance changed or the completed score is older than one day.",
        "A singleton lease prevents concurrent users from starting duplicate scoring jobs.",
        "The server sends a versioned feature payload to the configured XGBoost endpoint.",
        "Validated classifications, probabilities, anomaly scores and reasons are persisted as a versioned model run.",
        "The live dashboard refreshes, animates charts and makes the current filtered results exportable.",
    ])
    add_heading(document, "6.2 Safety and operation", 2)
    add_bullets(document, [
        "Model output is advisory and cannot directly alter attendance or payroll.",
        "Failures release the lease and set a 15-minute retry window.",
        "This is on-visit automation, not an unattended scheduler. A scheduled worker can be added later if required.",
        "The XGBoost URL and token remain server-only environment values.",
    ])

    page_break(document)
    add_heading(document, "7. Reports, exports, payslips and auditability", 1)
    add_table(document, ["Output", "Format", "Content / use"], [
        ("Operational records", "CSV and Excel", "Filtered employee, attendance, compensation, benefit, claim and payroll records"),
        ("Payroll register", "Excel", "Employee-level earnings, minutes, statutory deductions, contributions and net pay"),
        ("Payroll print package", "Print / PDF", "A4 payroll register followed by an individual payslip for each employee"),
        ("Analytics report", "CSV and Excel", "Current filtered live snapshot, classifications, anomalies and model metrics"),
        ("Audit history", "Application view / database", "Actor, action, entity, record, timestamp and retained deletion history"),
    ], [1.5, 1.2, 3.9])
    add_heading(document, "7.1 Payslip contents", 2)
    add_bullets(document, [
        "Payroll period, cutoff schedule, pay date, employee number and employee name.",
        "Basic pay, allowances, categorized overtime, night differential, bonus, benefits and reimbursements.",
        "Late/undertime/absence minutes with their deductions.",
        "SSS, PhilHealth, Pag-IBIG, BIR withholding, benefit and other deductions.",
        "Gross pay, total deductions, employer contributions and final net pay.",
    ])
    add_callout(document, "Privacy control", "Employees may only view their own payroll items through database row-level security. Administrative exports are limited to authorized payroll/HR roles.", INDIGO_LIGHT)

    page_break(document)
    add_heading(document, "8. Recommended operating procedure", 1)
    add_table(document, ["Timing", "Owner", "Activity", "Completion evidence"], [
        ("Ongoing", "HR admin", "Maintain employee status, salary history, benefits and approved leave", "Validated current records"),
        ("Daily / sync", "HR or integration", "Capture time-in/out, minutes and reviewed work-day type", "Attendance exceptions resolved"),
        ("Before cutoff", "HR/payroll", "Complete claims, bonuses and compensation approvals", "Approved inputs available"),
        ("Payday - 2 days", "Payroll manager", "Create/recalculate draft payroll and run validation", "No blocking issues"),
        ("Review window", "Authorized approver", "Compare totals, exceptions and sampled payslips", "Payroll approved and locked"),
        ("Payday", "Payroll/finance", "Mark paid after external disbursement; issue payslips", "Paid status and outputs retained"),
        ("Post-payroll", "HR/payroll", "Review analytics, audit changes and resolve anomalies", "Actions documented"),
    ], [1.1, 1.2, 3.0, 1.7])
    add_heading(document, "8.1 Exception handling", 2)
    add_table(document, ["Exception", "Required response"], [
        ("Missing compensation", "Create a non-overlapping effective salary record, then recalculate."),
        ("Attendance conflict", "Correct source record or reviewed classification; do not hide the exception with a manual net-pay edit."),
        ("Unapproved leave", "Complete the leave decision before calculation or treat the absence according to policy."),
        ("Failed payroll validation", "Resolve all blocking issues and failed historical comparisons before submission."),
        ("Statutory rule changed", "Create and approve a new effective-dated payroll policy; never rewrite a policy used by locked payroll."),
        ("AI service unavailable", "Use existing live records; investigate service configuration; automatic scoring retries after the cooldown."),
        ("Unauthorized action", "Confirm role, employment status, MFA assurance and department scope; never bypass RLS."),
    ], [2.0, 4.6])

    page_break(document)
    add_heading(document, "9. Deployment and acceptance checklist", 1)
    add_heading(document, "9.1 Database and authentication", 2)
    add_bullets(document, [
        "Apply migrations in documented order through 202609190001; never rerun already applied migrations.",
        "Configure Supabase URL, anonymous key and server-only service role key where needed.",
        "Create the protected first administrator, remove the one-time setup token and enroll two TOTP authenticators.",
        "Verify role assignment, manager scope, owner protection, inactive-account denial and employee self-access.",
    ])
    add_heading(document, "9.2 Payroll policy acceptance", 2)
    add_bullets(document, [
        "Obtain written approval for salary frequencies, cutoff allocation, workday hours and rounding.",
        "Confirm SSS, PhilHealth, Pag-IBIG and BIR configuration for the effective period.",
        "Confirm regular, rest-day, special-day, holiday, double-holiday and night premium treatment.",
        "Load anonymized historical comparison cases and reconcile representative payslips before live use.",
    ])
    add_heading(document, "9.3 Functional acceptance", 2)
    add_bullets(document, [
        "Create, edit, export and safely delete permitted draft records; confirm audit history remains.",
        "Test monthly-, daily- and hourly-paid workers across both cutoffs and a monthly run.",
        "Test late, undertime, absence, overtime, paid leave, benefit cost, claim reimbursement and bonus cases.",
        "Submit, approve, lock, print and export payroll; confirm locked entries cannot be changed.",
        "Run npm test, npm run typecheck, npm run lint and npm run build, followed by authenticated hosted-Supabase smoke checks.",
    ])
    add_callout(document, "Production sign-off", "Technical completion does not replace payroll/legal review. A Philippine payroll professional and the client's authorized finance/HR owners should approve the effective policy and reconciled results before the first live payroll.", RED_LIGHT)

    document.save(OUTPUT)


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    diagrams = [save_end_to_end(), save_access(), save_attendance(), save_payroll(), save_analytics()]
    build_document(diagrams)
    print(OUTPUT)


if __name__ == "__main__":
    main()
